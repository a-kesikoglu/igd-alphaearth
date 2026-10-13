"""Build docs/Extended_Results.docx (tables S1-S9, figures S1-S3, algorithms, archive description)."""
import json
from pathlib import Path
import docx, pandas as pd
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from common import ROOT

T = ROOT / "outputs/tables"; F = ROOT / "outputs/figures"; R = ROOT / "reports"
d = docx.Document()
for s in d.sections:
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2)
st = d.styles["Normal"]; st.font.name = "Cambria"; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria"); st.paragraph_format.space_after = Pt(4)


def para(text, bold=False, italic=False, size=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, mono=False):
    p = d.add_paragraph(); p.alignment = align
    r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if mono: r.font.name = "Consolas"; r.font.size = Pt(8.5)
    return p


def head(text):
    para(text, bold=True, size=11, align=WD_ALIGN_PARAGRAPH.LEFT)


def border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders")
    for e, v in kw.items():
        x = OxmlElement(f"w:{e}"); x.set(qn("w:val"), v); x.set(qn("w:sz"), "6"); x.set(qn("w:color"), "000000"); b.append(x)
    tcPr.append(b)


def table(csv, caption, note=None, size=8.5):
    p = para("", align=WD_ALIGN_PARAGRAPH.LEFT); a, b = caption.split(". ", 1)
    p.add_run(a + ". ").bold = True; p.add_run(b)
    df = pd.read_csv(T / csv, dtype=str, keep_default_na=False)
    t = d.add_table(rows=len(df) + 1, cols=len(df.columns))
    for j, c in enumerate(df.columns):
        cell = t.cell(0, j); cell.text = ""; r = cell.paragraphs[0].add_run(c); r.bold = True; r.font.size = Pt(size)
    for i, row in enumerate(df.itertuples(index=False), 1):
        for j, v in enumerate(row):
            cell = t.cell(i, j); cell.text = ""; r = cell.paragraphs[0].add_run(str(v)); r.font.size = Pt(size)
    for i, row in enumerate(t.rows):
        for c in row.cells:
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            kw = {}
            if i == 0: kw.update(top="single", bottom="single")
            if i == len(t.rows) - 1: kw.update(bottom="single")
            if kw: border(c, **kw)
    if note: para(note, size=8.5)
    para("")


def figure(png, caption, width=16):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(F / png), width=Cm(width))
    q = para("", align=WD_ALIGN_PARAGRAPH.CENTER); a, b = caption.split(". ", 1); q.add_run(a + ". ").bold = True; q.add_run(b)


ds = open(R / "dataset_audit.md", encoding="utf-8").read()
man = json.load(open(ROOT / "data/splits/manifest.json"))
para("Extended results and processing details", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.LEFT)
para("Sample Efficiency of AlphaEarth Embeddings for Land Cover Classification: Spatial Holdout and Cross-Country Evaluation",
     italic=True, align=WD_ALIGN_PARAGRAPH.LEFT)

head("S1. Reference data selection and feature extraction")
para("Table S1 gives the number of LUCAS 2018 Copernicus polygons after each selection step, counted in Google Earth Engine. "
     "All features were extracted with the Earth Engine reduceRegions operation and a mean reducer at a nominal scale of 10 m. "
     "Because the mosaicked collections carry no fixed projection, the reduction was evaluated on a 10 m grid in geographic "
     "coordinates (EPSG:4326); every grid cell took the value of the native 10 m pixel at its location (nearest neighbour), and "
     "cells were weighted by the fraction of their area inside the polygon. For polygon 22543 (1,566 m²), for example, 20 grid "
     "cells had their centre inside the polygon and the summed fractional weights were 20.47. The mean therefore approximates an "
     "area weighted mean of the native pixels. The polygons contained 3 to 77 grid cell centres (median 43).")
para("Sentinel-2 surface reflectance scenes of 2018 were linked to Cloud Score+ by scene identifier and masked where the "
     "cumulative distribution score (cs_cdf) was below 0.60. For every polygon and quarter, the area fraction with at least one "
     "clear observation was computed with the same reducer; it equalled 1.000 for all 10,144 polygons in all four quarters, so no "
     "polygon had an empty composite. Sentinel-1 ground range detected scenes (interferometric wide swath, VV and VH) were used as "
     "distributed in Earth Engine, that is, after thermal noise removal, radiometric calibration to backscatter coefficient and "
     "orthorectification with a digital elevation model, without radiometric terrain flattening, and expressed in decibels. "
     "Quarterly medians were computed per pixel from both orbit directions. Table S9 lists the scene counts; per grid cell, "
     "every polygon had at least 6 ascending and 11 descending observations in every quarter. Medians from a single orbit "
     "direction were evaluated as a sensitivity analysis (Table S5).")
table("tableS1_selection.csv", "Table S1. Sequential selection of LUCAS 2018 Copernicus polygons")

head("S2. Algorithms")
para("Test block cluster bootstrap with resampled training draws (used for all intervals in the main text):", align=WD_ALIGN_PARAGRAPH.LEFT)
for line in [
    "for each test set (Spain: G = 62 blocks of 50 km; Greece: G = 112 blocks of 50 km):",
    "  C[r, g] = confusion matrix of repetition r restricted to test block g",
    "  draw block weights w*[b, g] ~ Multinomial(G, 1/G), b = 1..2000   (shared by all models and budgets)",
    "  for each budget n: draw repetition weights v*[b, r] ~ Multinomial(R, 1/R)",
    "      (shared by all representations and classifiers of budget n; independent between budgets)",
    "  for each representation and classifier at budget n:",
    "      F*[b, r] = macro F1 of sum_g w*[b, g] C[r, g]",
    "      F*[b]    = sum_r v*[b, r] F*[b, r] / R",
    "  difference A - B: D*[b] = F*_A[b] - F*_B[b];  interval = 2.5th and 97.5th percentiles of D*"]:
    para(line, mono=True, align=WD_ALIGN_PARAGRAPH.LEFT)
para("Intervals conditional on the training draws (block weights only, v* = 1) are archived with the results (columns ci_lo and "
     "ci_hi) and are on average about 25% narrower.")
para("Dissimilarity index (Equation 9 of the main text):", align=WD_ALIGN_PARAGRAPH.LEFT)
for line in [
    "standardise every feature with the mean and SD of the Spanish training pool (6,129 polygons)",
    "d_train = median over pool samples of the Euclidean distance to the nearest OTHER pool sample (self excluded)",
    "for each test sample j: DI_j = Euclidean distance to the nearest pool sample / d_train",
    "no feature weighting is applied; one nearest neighbour is used"]:
    para(line, mono=True, align=WD_ALIGN_PARAGRAPH.LEFT)
para("The index follows the idea of the area of applicability of Meyer and Pebesma (2021) but omits their weighting of "
     "predictors by variable importance and their threshold derived from cross-validation folds.")

head("S3. Additional results")
table("tableS2_rf_all.csv", "Table S2. Macro F1 score of Random Forest for all feature sets (mean ± standard deviation over 20 training repetitions; five seeds for the full pool)")
table("tableS3_lr_all.csv", "Table S3. Macro F1 score of multinomial logistic regression for all feature sets (mean ± standard deviation over 20 training repetitions; one fit for the full pool)",
      note="Note: Logistic regression is deterministic, so the full pool row has no variation.")
table("tableS4_paired_delta.csv", "Table S4. Paired macro F1 differences with 95% two-level bootstrap intervals")
table("tableS5_sensitivity.csv", "Table S5. Sensitivity analyses with Random Forest: contribution of Sentinel-1, single orbit direction and unit-length embeddings (paired differences with 95% two-level bootstrap intervals)")
table("tableS6_cross_budget.csv", "Table S6. Comparisons across label budgets (paired differences with 95% two-level bootstrap intervals; repetitions of different budgets resampled independently)")
table("tableS7_ladder.csv", "Table S7. Test and training set sizes and distances in the validation sensitivity analysis (Random Forest, median over ten repetitions)",
      note="Note: The sensitivity analysis uses no buffer and redraws the held out samples or blocks in every repetition; it is therefore not identical to the locked holdout.")
table("tableS8_dissimilarity.csv", "Table S8. Dissimilarity of test samples to the Spanish training pool")
table("tableS9_s1_scenes.csv", "Table S9. Sentinel-1 interferometric wide swath scenes with VV and VH polarisation intersecting the study region in 2018, by quarter and orbit direction")

head("S4. Additional figures")
figure("fig2_workflow.png", "Figure S1. Workflow of data preparation, training and evaluation.", 15)
figure("figS2_confusion.png", "Figure S2. Row normalised confusion matrices (%) of Random Forest with 100 labels per class, pooled over the 20 training repetitions (the pooled counts are not independent samples). Supports: Spain 749, 485, 305, 376, 134; Greece 511, 468, 205, 545, 56 (cropland, woodland, shrubland, grassland, bare land). Cr, cropland; Wo, woodland; Sh, shrubland; Gr, grassland; Ba, bare land.", 16)
figure("figS3_dissimilarity.png", "Figure S3. Cumulative distributions of the dissimilarity index of Spanish test and Greek polygons for three representations; the vertical line marks DI = 1.", 16)

head("S5. Data and code archive")
para("The archive contains the configuration, all scripts and the following files: the reference and feature tables extracted from "
     "Earth Engine (data/raw), the analysis table (data/processed/samples.parquet), the split identifier lists with SHA-256 "
     "checksums and the 50 km test block geometries (data/splits), per run metrics (outputs/metrics/e1_runs.csv, e2_ladder.csv), "
     "a long format prediction table with one row per test sample and run (outputs/metrics/predictions_long.parquet; columns "
     "run_id, test_set, sample_id, point_id, block_id, true_label, prediction, representation, classifier, budget, repetition, "
     "rf_seed, sampling_seed), and the bootstrap outputs. The command python reproduce_all.py regenerates all tables and figures "
     "from the archived feature tables; python reproduce_all.py --extract repeats the Earth Engine extraction. Software: "
     "Python 3.12.10, scikit-learn 1.9.1, earthengine-api 1.7.46 (full lock file in reports/requirements-lock.txt).")
para("Split checksums: " + "; ".join(f"{k}: {v['n']} polygons, SHA-256 {v['sha256'][:16]}…" for k, v in man["files"].items()) + ".", size=8.5)
(ROOT / "docs").mkdir(exist_ok=True); out = ROOT / "docs/Extended_Results.docx"; d.save(out); print("saved", out)
